"""Regras de contato compartilhadas entre a geracao via banco e via Excel.

Todo disparo termina em CSVs de duas colunas (phonenumber, name), separados
por grupo de rating. Este modulo concentra normalizacao, deduplicacao,
classificacao por rating e escrita dos arquivos.
"""

from __future__ import annotations

import csv
import shutil
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Iterable, NamedTuple

from openpyxl import load_workbook

# Cada grupo vira um CSV de saida. O contencioso nao e separado por rating.
GROUP_ABW = "abw"
GROUP_C = "c"
GROUP_DEZ = "dez"
GROUP_NA_RATING = "sem_rating"
GROUP_CONTENCIOSO = "contencioso"

OUTPUT_FILES = {
    GROUP_ABW: "amigavel_ABW.csv",
    GROUP_C: "amigavel_C.csv",
    GROUP_DEZ: "amigavel_DEZ.csv",
    GROUP_NA_RATING: "amigavel_na_rating.csv",
    GROUP_CONTENCIOSO: "contencioso.csv",
}

# Caminho inverso, csv -> grupo. Usado pela tela pra ligar o dispatches.json
# (que so conhece o nome do csv) de volta ao grupo, sem duplicar as chaves.
CSV_PARA_GRUPO = {arquivo: grupo for grupo, arquivo in OUTPUT_FILES.items()}

# O rating vem como "A", "B", "C", "D", "E" ou prefixado ("Z_REDUCAO",
# "W_FPD_COM_PL", ...). A primeira letra define o grupo.
RATING_GROUPS = {
    "A": GROUP_ABW,
    "B": GROUP_ABW,
    "W": GROUP_ABW,
    "C": GROUP_C,
    "D": GROUP_DEZ,
    "E": GROUP_DEZ,
    "Z": GROUP_DEZ,
}

TIPO_AMIGAVEL = "amigavel"
TIPO_CONTENCIOSO = "contencioso"
TIPOS_VALIDOS = (TIPO_AMIGAVEL, TIPO_CONTENCIOSO)

# Telefone brasileiro valido tem no minimo DDD + 8 digitos.
MIN_PHONE_DIGITS = 10


class Registro(NamedTuple):
    """Linha bruta, antes de qualquer normalizacao."""

    telefone: object
    nome: object
    tipo: object
    rating: object


@dataclass
class ResultadoContatos:
    grupos: dict[str, list[tuple[str, str]]]
    total_lidos: int = 0
    telefone_invalido: int = 0
    sem_nome: int = 0
    duplicados: int = 0
    ratings_desconhecidos: Counter = field(default_factory=Counter)

    @property
    def total_contatos(self) -> int:
        return sum(len(rows) for rows in self.grupos.values())


def normalize_phone(value: object) -> str | None:
    """Deixa apenas digitos e descarta o que for curto demais para disparo."""
    if value is None:
        return None

    if isinstance(value, int):
        digits = str(value)
    elif isinstance(value, float):
        digits = str(int(value)) if value.is_integer() else "".join(ch for ch in str(value) if ch.isdigit())
    else:
        digits = "".join(ch for ch in str(value).strip() if ch.isdigit())

    if len(digits) < MIN_PHONE_DIGITS:
        return None

    return digits


def normalize_name(value: object) -> str | None:
    if value is None:
        return None

    name = " ".join(str(value).split())
    if not name:
        return None

    return name.lower().title()


def normalize_header(value: object) -> str:
    return str(value).strip().lower()


def normalize_sheet_name(name: object) -> str:
    return str(name).strip().lower()


def eh_contencioso(tipo: object) -> bool:
    return tipo is not None and TIPO_CONTENCIOSO in str(tipo).strip().lower()


def rating_group(rating: object) -> str | None:
    """Grupo do rating amigavel. None quando vazio ou desconhecido."""
    if rating is None:
        return None

    rating_text = str(rating).strip().upper()
    if not rating_text:
        return None

    return RATING_GROUPS.get(rating_text[0])


def rating_desconhecido(rating: object) -> bool:
    """True quando o rating veio preenchido mas nao bate com nenhum grupo."""
    if rating is None:
        return False

    rating_text = str(rating).strip().upper()
    return bool(rating_text) and rating_text[0] not in RATING_GROUPS


def resolve_group(tipo: object, rating: object) -> str:
    """Define em qual CSV a linha entra.

    Contencioso ignora o rating. Amigavel usa a primeira letra do rating;
    sem rating (ou rating desconhecido) vai para o CSV proprio.
    """
    if eh_contencioso(tipo):
        return GROUP_CONTENCIOSO

    return rating_group(rating) or GROUP_NA_RATING


def coletar_contatos(registros: Iterable[Registro]) -> ResultadoContatos:
    """Normaliza, deduplica por telefone e agrupa por rating.

    A deduplicacao e global: o mesmo telefone nao pode receber dois disparos,
    mesmo que apareca em grupos diferentes.
    """
    resultado = ResultadoContatos(grupos={group: [] for group in OUTPUT_FILES})
    vistos: set[str] = set()

    for registro in registros:
        resultado.total_lidos += 1

        phone = normalize_phone(registro.telefone)
        if not phone:
            resultado.telefone_invalido += 1
            continue

        name = normalize_name(registro.nome)
        if not name:
            resultado.sem_nome += 1
            continue

        if phone in vistos:
            resultado.duplicados += 1
            continue

        if not eh_contencioso(registro.tipo) and rating_desconhecido(registro.rating):
            resultado.ratings_desconhecidos[str(registro.rating).strip().upper()] += 1

        vistos.add(phone)
        resultado.grupos[resolve_group(registro.tipo, registro.rating)].append((phone, name))

    return resultado


# Extensoes aceitas na pasta de filtro. csv entra pra facilitar exportacao rapida.
FILTER_FILE_PATTERNS = ("*.xlsx", "*.xlsm", "*.xltx", "*.xltm", "*.csv")


def coletar_arquivos_filtro(pasta: Path) -> list[Path]:
    if not pasta.exists():
        return []

    arquivos: list[Path] = []
    for pattern in FILTER_FILE_PATTERNS:
        arquivos.extend(pasta.glob(pattern))

    return sorted(arquivos)


def _telefones_do_csv(arquivo: Path) -> Iterable[str]:
    with arquivo.open(encoding="utf-8-sig", newline="") as f:
        for linha in csv.reader(f):
            for valor in linha:
                telefone = normalize_phone(valor)
                if telefone:
                    yield telefone


def _telefones_do_excel(arquivo: Path) -> Iterable[str]:
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for linha in sheet.iter_rows(values_only=True):
                for valor in linha:
                    telefone = normalize_phone(valor)
                    if telefone:
                        yield telefone
    finally:
        workbook.close()


def ler_telefones_filtro(pasta: Path) -> set[str]:
    """Le toda planilha (xlsx ou csv) da pasta de filtro e devolve os telefones a remover.

    Nao exige cabecalho nem coluna fixa: qualquer celula que normalize para um
    telefone valido entra no filtro, pra aceitar listas coladas sem formatacao.
    """
    telefones: set[str] = set()

    for arquivo in coletar_arquivos_filtro(pasta):
        extrator = _telefones_do_csv if arquivo.suffix.lower() == ".csv" else _telefones_do_excel
        telefones.update(extrator(arquivo))

    return telefones


def aplicar_filtro(
    grupos: dict[str, list[tuple[str, str]]], telefones_filtro: set[str]
) -> tuple[dict[str, list[tuple[str, str]]], dict[str, list[tuple[str, str]]]]:
    """Remove das bases os telefones presentes no filtro.

    Devolve (grupos_filtrados, removidos_por_grupo) — removidos_por_grupo guarda
    as linhas (telefone, nome) removidas de cada grupo, pra auditoria e pro
    detalhamento na tela (antes so devolvia um total agregado).
    """
    if not telefones_filtro:
        return grupos, {}

    removidos_por_grupo: dict[str, list[tuple[str, str]]] = {}
    filtrados: dict[str, list[tuple[str, str]]] = {}

    for grupo, linhas in grupos.items():
        mantidos = [linha for linha in linhas if linha[0] not in telefones_filtro]
        removidos = [linha for linha in linhas if linha[0] in telefones_filtro]
        if removidos:
            removidos_por_grupo[grupo] = removidos
        filtrados[grupo] = mantidos

    return filtrados, removidos_por_grupo


def escrever_filtro_removidos(
    pasta: Path, data_referencia: date, removidos_por_grupo: dict[str, list[tuple[str, str]]]
) -> Path:
    """Grava os numeros removidos pelo filtro manual, pra conferencia.

    Mesmo padrao de nome incremental do gravar_relatorio (nao sobrescreve
    rodadas do mesmo dia).
    """
    pasta.mkdir(parents=True, exist_ok=True)
    nome = f"filtro_removidos_{data_referencia:%Y%m%d}"
    arquivo = pasta / f"{nome}.csv"

    contador = 1
    while arquivo.exists():
        contador += 1
        arquivo = pasta / f"{nome}_{contador}.csv"

    with arquivo.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["telefone", "nome", "grupo"])
        for grupo, linhas in removidos_por_grupo.items():
            for telefone, nome_contato in linhas:
                writer.writerow([telefone, nome_contato, grupo])

    return arquivo


def clear_output_folder(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    for item in output_dir.iterdir():
        if item.name == ".gitkeep" and item.is_file():
            continue

        if item.is_dir():
            shutil.rmtree(item)
        else:
            item.unlink()


def write_csv(file_path: Path, data: Iterable[tuple[str, str]]) -> None:
    with file_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["phonenumber", "name"])
        writer.writerows(data)


def escrever_grupos(output_dir: Path, grupos: dict[str, list[tuple[str, str]]]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    for group, file_name in OUTPUT_FILES.items():
        write_csv(output_dir / file_name, grupos.get(group, []))


def imprimir_resumo(resultado: ResultadoContatos) -> None:
    print("Arquivos gerados:")
    for group, file_name in OUTPUT_FILES.items():
        print(f"- {file_name} ({len(resultado.grupos.get(group, []))} contatos)")

    print(f"Total pronto para disparo: {resultado.total_contatos} contatos.")
    print(
        f"Descartados: {resultado.telefone_invalido} telefone invalido, "
        f"{resultado.sem_nome} sem nome, {resultado.duplicados} duplicados "
        f"(de {resultado.total_lidos} linhas lidas)."
    )

    if resultado.ratings_desconhecidos:
        detalhes = ", ".join(
            f"{rating} ({quantidade})"
            for rating, quantidade in sorted(resultado.ratings_desconhecidos.items())
        )
        print(f"AVISO: ratings desconhecidos enviados para o CSV sem rating: {detalhes}.")
